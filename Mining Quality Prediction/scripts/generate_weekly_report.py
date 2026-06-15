"""Generate Weekly Progress Report as Word document."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUTPUT = Path(__file__).resolve().parent.parent / "Weekly_Progress_Report_MinePredict_AI.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row_idx, row in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)
    doc.add_paragraph()


def main():
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Weekly Progress Report")
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph()
    meta = [
        ("Name:", "___________________________"),
        ("Domain:", "Machine Learning / Full-Stack Web Development (Mining Process Optimization)"),
        ("Date of submission:", "May 20, 2026"),
        ("Week Ending:", "Week 01"),
        ("Project name:", "Interactive Mining Quality Predictor (MinePredict AI / MineVision AI)"),
        ("Reference basis:", "ML-based silica concentrate prediction in froth flotation (industrial dataset, March–September 2017)"),
        ("Overall project completion:", "25%"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(label + " ").bold = True
        p.add_run(value)

    doc.add_paragraph()

    add_heading(doc, "I. Overview", 1)
    doc.add_paragraph(
        "This week focused on establishing the foundation of an AI-powered mining quality prediction "
        "platform. The goal is to predict % Silica Concentrate in iron ore flotation plants using "
        "machine learning, replacing slow laboratory tests (typically 2+ hours) with near real-time "
        "predictions."
    )
    doc.add_paragraph(
        "Work covered roughly one quarter (25%) of the full project scope: core frontend shell "
        "(Home, Dashboard, Login), FastAPI backend with ML pipeline, initial ensemble model training "
        "on the flotation plant dataset, and basic API integration between frontend and backend."
    )
    doc.add_paragraph(
        "Remaining work (~75%) includes advanced UI sections, multi-horizon forecasting improvements, "
        "dataset analytics pages, mining sites module, deployment, and enterprise features (chatbot, "
        "3D visuals, etc.)."
    )

    add_heading(doc, "II. Achievements", 1)

    add_heading(doc, "1. Project Setup & Architecture (Phase 1 — Complete)", 2)
    for item in [
        "Initialized React 18 + TypeScript + Vite frontend with Tailwind CSS v4 and Radix UI components.",
        "Built FastAPI backend with modular routers: Authentication, Prediction, Training, Dataset insights, Dashboard metrics.",
        "Configured Vite proxy so frontend calls /api → backend on port 8000.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "2. Frontend Development (~30% of UI scope)", 2)
    add_table(
        doc,
        ["Module", "Status", "Notes"],
        [
            ("Home Page", "Partial", "Hero, process flow, model overview, mining theme"),
            ("Login", "Basic", "Employee/Admin JWT login"),
            ("Dashboard", "Partial", "Predict / Train / Analytics tabs"),
            ("Dataset Uploader", "Basic", "Drag-and-drop CSV upload"),
            ("Navigation", "Minimal", "Only Home, Dashboard, Login"),
            ("Mining Sites", "Not started", "—"),
            ("Dataset Insights Page", "Not started", "—"),
            ("About / Contact", "Not started", "Footer placeholders only"),
        ],
    )

    add_heading(doc, "3. Backend & ML Model (~40% of backend scope)", 2)
    for item in [
        "Extracted and loaded the Mining Process Flotation Plant Database (March–September 2017, 22 columns).",
        "Implemented preprocessing: hourly alignment, lag features, rolling statistics, engineered features (37 total).",
        "Trained ensemble model (Random Forest + XGBoost + LightGBM).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Current model performance (validation set):")
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ("R²", "0.56"),
            ("RMSE", "0.76% silica"),
            ("MAE", "0.58% silica"),
            ("Training samples", "3,276"),
            ("Validation samples", "820"),
        ],
    )
    doc.add_paragraph(
        "Multi-horizon models exist (1h / 3h / 6h) but longer horizons need improvement "
        "(3h R² ≈ 0.03, 6h R² negative)."
    )

    add_heading(doc, "4. Real-Time Prediction Interface (~25%)", 2)
    for item in [
        "Dashboard input form for process parameters (iron feed, silica feed, starch/amina flow, pulp pH/density, flotation columns).",
        "API returns predicted silica %, risk level, confidence, recommendations, and horizon forecasts.",
        "Charts use Recharts; some analytics data is still mock/placeholder.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "5. Authentication (~50%)", 2)
    for item in [
        "JWT-based login for Employee and Admin roles.",
        "Demo credentials: employee@minevision.ai / employee123.",
        "Token stored in localStorage; protected API calls implemented.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "III. Challenges", 1)

    add_heading(doc, "1. Model Accuracy vs. Target", 2)
    for item in [
        "README/marketing claims ~94% accuracy, but validated test R² is ~56%.",
        "Long-horizon prediction (3h, 6h) is weak and needs LSTM/time-series work (planned but not implemented).",
        "Gap between train metrics (R² ≈ 0.98) and test metrics suggests possible overfitting; regularization and cross-validation are needed.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "2. Incomplete Frontend Scope", 2)
    doc.add_paragraph(
        "Full specification includes 12 major sections; only ~3 are partially built. Missing: Mining Sites, "
        "Dataset Insights page, model comparison UI, AI chatbot, 3D tunnel walkthrough, full navigation."
    )

    add_heading(doc, "3. Frontend–Backend Integration Gaps", 2)
    for item in [
        "Dashboard still uses mock historical data when API is offline.",
        "TensorFlow.js is imported but not fully wired for client-side inference.",
        "Dataset insights endpoint exists but no dedicated UI page yet.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "4. Data & Environment", 2)
    for item in [
        "Dataset uses European decimal format (comma separators); preprocessing handles this but edge cases may remain.",
        "Backend requires manual dataset extraction via extract_dataset.py on first run.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "IV. Learning Resources", 1)

    add_heading(doc, "1. Domain & Research", 2)
    doc.add_paragraph(
        "Thesis: Machine Learning-based Quality Prediction in the Froth Flotation Process of Mining "
        "(Eric Kwame Osei, Dalarna University, 2019)."
    )
    doc.add_paragraph(
        "Dataset: Google Drive – Mining Process Flotation Plant Database (March–September 2017)."
    )

    add_heading(doc, "2. Technical Stack", 2)
    doc.add_paragraph(
        "Frontend: React Router v7, Recharts, Motion (Framer Motion), TensorFlow.js, Tailwind CSS v4, Radix UI, Vite."
    )
    doc.add_paragraph(
        "Backend: FastAPI, scikit-learn, XGBoost, LightGBM, Pandas, NumPy."
    )

    add_heading(doc, "3. Tools Used", 2)
    for item in [
        "Cursor IDE for development",
        "FastAPI Swagger docs (http://localhost:8000/docs) for API testing",
        "pnpm for frontend dependencies",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "V. Next Week's Goals", 1)
    doc.add_paragraph("(Aligned with remaining ~75% of project scope)")

    add_heading(doc, "1. Model Enhancement (Priority: High)", 2)
    for item in [
        "Improve test R² toward ≥ 0.75 via hyperparameter tuning and cross-validation.",
        "Implement LSTM for multi-step ahead prediction (1h, 3h, 6h).",
        "Add model comparison UI (RF vs XGBoost vs LightGBM vs Neural Network).",
        "Reduce overfitting (regularization, feature selection).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "2. Frontend Expansion (Priority: High)", 2)
    for item in [
        "Add routes: Dataset Insights, Model Training, Mining Sites, About, Contact.",
        "Replace mock dashboard data with live API data.",
        "Build correlation matrix, feature importance charts, and preprocessing pipeline visualization.",
        "Expand input form to all 22 dataset columns with tooltips and sliders.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "3. Backend & Integration (Priority: Medium)", 2)
    for item in [
        "Complete dataset insights API consumption in frontend.",
        "Add CSV batch prediction endpoint.",
        "Improve horizon model accuracy and confidence scoring.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "4. UI/UX & Advanced Features (Priority: Lower — later phases)", 2)
    for item in [
        "Full navigation bar per specification.",
        "AI chatbot assistant for engineers.",
        "3D mining tunnel visualization (Three.js).",
        "Dark/light industrial theme toggle.",
        "Docker deployment and cloud integration documentation.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "VI. Additional Comments", 1)
    doc.add_paragraph(
        "The project has a solid technical foundation at ~25% completion: working frontend shell, "
        "functional backend API, trained ensemble model, and basic prediction flow. The core problem — "
        "real-time silica concentrate prediction in flotation plants — is correctly scoped and aligned "
        "with industrial research."
    )

    add_heading(doc, "Strengths at this stage:", 2)
    for item in [
        "End-to-end path from login → dashboard → prediction API",
        "Real industrial dataset integrated",
        "Mining-themed UI with process flow visualization",
        "Modular, extensible backend architecture",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "Gaps to address before milestone 2 (~50%):", 2)
    for item in [
        "Model accuracy validation and honest metrics in UI",
        "Dedicated analytics and dataset insights pages",
        "Removal of mock data dependencies",
        "Broader navigation and missing major sections",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Estimated timeline (assuming continued weekly effort):")
    add_table(
        doc,
        ["Milestone", "Target completion", "Focus"],
        [
            ("25% (current)", "Week 1", "Foundation, basic predict/train"),
            ("50%", "Week 3–4", "Dataset insights, model tuning, full dashboard"),
            ("75%", "Week 6–7", "Mining sites, LSTM, advanced analytics"),
            ("100%", "Week 8–10", "Chatbot, 3D visuals, deployment, polish"),
        ],
    )

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run("Prepared for: ").bold = True
    footer.add_run("Learning Project / Minor Project submission")
    doc.add_paragraph("Project repository: Interactive Mining Quality Predictor")
    doc.add_paragraph("Report format source: Weekly Report format UPDATED.pdf")

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
